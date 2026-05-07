from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "marketing-os-partner-proposal.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa: int) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")


def set_cell_margins(table, margin_dxa: int = 110) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.first_child_found_in("w:tblCellMar")
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for side in ("top", "left", "bottom", "right"):
        node = tbl_cell_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(margin_dxa))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(3)


def add_feature_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_width(table, 9360)
    set_cell_margins(table, 120)

    widths = [2700, 6660]
    header = table.rows[0]
    set_repeat_table_header(header)
    header.cells[0].text = "Feature"
    header.cells[1].text = "What It Provides"
    for i, cell in enumerate(header.cells):
        set_cell_width(cell, widths[i])
        set_cell_shading(cell, "EDEDED")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    for feature, description in rows:
        row = table.add_row()
        row.cells[0].text = feature
        row.cells[1].text = description
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    doc.add_paragraph()


def add_api_table(doc: Document, rows: list[tuple[str, str, str]]) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_width(table, 9360)
    set_cell_margins(table, 100)

    widths = [2100, 3400, 3860]
    header = table.rows[0]
    set_repeat_table_header(header)
    labels = ["API Group", "Example Endpoints", "Purpose"]
    for i, label in enumerate(labels):
        cell = header.cells[i]
        cell.text = label
        set_cell_width(cell, widths[i])
        set_cell_shading(cell, "EDEDED")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    for group, endpoints, purpose in rows:
        row = table.add_row()
        values = [group, endpoints, purpose]
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.text = value
            set_cell_width(cell, widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    doc.add_paragraph()


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(22)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(13)

    header = section.header.paragraphs[0]
    header.text = "Marketing OS Partner Proposal"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer = section.footer.paragraphs[0]
    footer.text = "Editable Word document"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Marketing OS Partner Proposal").bold = True

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Description, API Features, and Full Marketing OS System Features")

    doc.add_heading("1. Description", level=1)
    doc.add_paragraph(
        "Marketing OS is a partner-ready marketing and messaging platform. A SaaS partner can onboard customers "
        "from their own software, create a customer workspace in Marketing OS, connect each customer's WhatsApp "
        "Business number, and send notifications or campaigns from each customer's own WhatsApp number."
    )
    doc.add_paragraph(
        "The partner can use Marketing OS in two ways. First, they can use APIs inside their own software. "
        "Second, they can give the customer access to the full Marketing OS system, where the customer can manage "
        "templates, broadcasts, campaigns, WhatsApp leads, tasks, automations, Instagram automation, analytics, "
        "catalog, and flows."
    )

    doc.add_heading("How Customer WhatsApp Details Are Created", level=2)
    doc.add_paragraph(
        "When a customer is onboarded through the partner system, the partner first creates the customer tenant in "
        "Marketing OS. The customer then connects WhatsApp using Meta Embedded Signup. During this signup, the customer "
        "selects their Meta Business Portfolio, WhatsApp Business Account, and WhatsApp phone number."
    )
    for item in [
        "After signup, Meta returns an authorization code to Marketing OS.",
        "Marketing OS exchanges that code with Meta and receives the customer's WABA ID, phone number ID, display phone number, business name, and access token/system token.",
        "Marketing OS stores those credentials securely against the customer's tenant/workspace.",
        "The partner can retrieve safe connection details such as tenant ID, WABA ID, phone number ID, display phone number, business name, and connection status.",
        "The raw Meta access token should normally stay inside Marketing OS and should not be exposed to the partner or customer dashboard.",
        "When the partner sends a message using x-tenant-id, Marketing OS loads that customer's stored phone number ID and token internally and sends from that customer's WhatsApp number.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("2. API Section: Features Partner Software Can Use", level=1)
    doc.add_paragraph(
        "This section is for partner software integration. These features can be used by the partner's SaaS through API, "
        "so the partner can add WhatsApp, Instagram, automation, and lead operations without building everything from scratch."
    )

    add_api_table(
        doc,
        [
            (
                "Customer onboarding",
                "POST /api/v1/tenants\nGET /api/v1/tenants\nPOST /api/v1/tenants/:tenantId/token",
                "Create customer workspaces, list customers, and generate customer dashboard access tokens.",
            ),
            (
                "Customer mapping",
                "GET /api/v1/customer-engine-maps\nPOST /api/v1/customer-engine-maps\nPUT /api/v1/customer-engine-maps/:id",
                "Map partner customer IDs to Marketing OS tenants and connected engines.",
            ),
            (
                "WhatsApp notifications",
                "POST /api/v1/messages/send\nGET /api/v1/messages\nGET /api/v1/messages/:messageId/status",
                "Send notifications from each client's connected WhatsApp number and track message status.",
            ),
            (
                "WhatsApp account details",
                "POST /api/v1/tenants/:tenantId/whatsapp\nGET /api/v1/tenants/:tenantId/whatsapp",
                "Link or read safe WhatsApp account details such as WABA ID, phone number ID, display number, business name, and connection status. Raw tokens should remain secured inside Marketing OS.",
            ),
            (
                "Template CRUD",
                "GET /api/v1/whatsapp/templates\nPOST /api/v1/whatsapp/templates\nPUT /api/v1/whatsapp/templates/:id\nPOST /api/v1/whatsapp/templates/:id/submit\nDELETE /api/v1/whatsapp/templates/:id",
                "Create, list, view, update, submit to Meta, sync, test, and delete WhatsApp templates.",
            ),
            (
                "Campaign / broadcasting",
                "POST /api/v1/whatsapp/broadcast\nGET /api/v1/whatsapp/broadcast\nGET /api/v1/whatsapp/analytics/campaigns",
                "Send broadcast campaigns, list broadcasts, view results, and track campaign analytics.",
            ),
            (
                "Flow building",
                "GET /api/v1/whatsapp/flows\nPOST /api/v1/whatsapp/flows\nPUT /api/v1/whatsapp/flows/:flowId\nPOST /api/v1/whatsapp/flows/:flowId/publish",
                "Create, update, publish, sync, and delete WhatsApp flows.",
            ),
            (
                "Automation",
                "POST /api/v1/events\nPOST /api/v1/events/mappings\nGET /api/v1/events/mappings",
                "Trigger automation from partner events such as lead created, booking confirmed, payment received, or order updated.",
            ),
            (
                "WhatsApp inbound leads",
                "GET /api/v1/whatsapp/conversations\nGET /api/v1/whatsapp/conversations/:id/messages\nPOST /api/v1/whatsapp/conversations/:id/send\nPOST /api/v1/whatsapp/conversations/:id/assign",
                "Handle inbound WhatsApp messages, replies, lead conversations, assignments, notes, and escalation.",
            ),
            (
                "Instagram automation",
                "GET /api/v1/instagram/inbox/comments\nPOST /api/v1/instagram/inbox/comments/:accountId/:commentId/private-reply\nGET /api/v1/instagram/inbox/messages\nPOST /api/v1/instagram/inbox/messages/:accountId/send",
                "Manage Instagram comments, comment to DM, DM automation, inbox, and lead capture.",
            ),
            (
                "Analytics",
                "GET /api/v1/whatsapp/analytics/campaigns\nGET /api/v1/whatsapp/analytics/response-time\nGET /api/v1/instagram/analytics/:accountId",
                "Track campaign performance, response time, Instagram insights, and lead performance.",
            ),
        ],
    )

    doc.add_heading("Important API Feature List", level=2)
    for item in [
        "Partner can onboard customers through their own system.",
        "Each customer gets a separate Marketing OS tenant/workspace.",
        "Each customer can connect and use their own WhatsApp number.",
        "Partner can manage notifications from each client WhatsApp number.",
        "Partner can send messages, trigger events, and track delivery status.",
        "Partner can use template CRUD, campaign/broadcasting, flow building, automation, lead management, Instagram automation, and analytics APIs.",
    ]:
        add_bullet(doc, item)

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_heading("3. Full System Section: Features Customer Gets Inside Marketing OS", level=1)
    doc.add_paragraph(
        "This section is for customers who use the full Marketing OS system after being onboarded by the partner. "
        "The customer can log in and manage their own marketing operations directly."
    )

    add_feature_table(
        doc,
        [
            ("WhatsApp workspace", "Connect and manage their own WhatsApp Business number, connection settings, and messaging setup."),
            ("Notifications from own number", "Send booking updates, payment reminders, confirmations, alerts, and follow-ups from their own WhatsApp number."),
            ("Template management", "Create, edit, test, submit, sync, and manage WhatsApp templates."),
            ("Campaign management", "Create and manage marketing campaigns for customer lists, segments, and follow-ups."),
            ("Broadcasting", "Send bulk WhatsApp broadcasts using approved templates and track campaign results."),
            ("Video/photo campaigns", "Use media campaign content where WhatsApp, Instagram, and Meta channel rules allow it."),
            ("Photo carousel", "Create Instagram carousel campaigns with up to 10 media items where supported by Instagram permissions."),
            ("WhatsApp lead management", "Capture inbound WhatsApp messages as leads, reply to conversations, assign team members, add notes, and follow up."),
            ("Lead and task management", "Create tasks, follow-ups, reminders, and internal work items for sales/support teams."),
            ("Advanced flow builder", "Build structured WhatsApp flows for lead qualification, forms, booking steps, service requests, and guided journeys."),
            ("Automation builder", "Build automated replies, routing, follow-ups, lead assignment, task creation, and event-based workflows."),
            ("Instagram automation", "Automate Instagram comments, comment-to-DM flows, DM replies, lead capture, and follow-up journeys."),
            ("Meta lead capture", "Capture leads from Meta/Instagram sources and manage them inside Marketing OS CRM."),
            ("Lead analytics", "View lead source, lead status, campaign response, conversion tracking, and performance reports."),
            ("CRM and inbox", "Manage contacts, conversations, tickets, notes, assignments, timelines, and escalation."),
            ("Catalog management", "Manage product/catalog content for WhatsApp commerce/catalog workflows."),
        ],
    )

    doc.add_heading("Customer Benefits", level=2)
    for item in [
        "Customer gets a complete marketing system, not only API sending.",
        "Customer can manage WhatsApp, Instagram, leads, tasks, automations, campaigns, and analytics in one place.",
        "Customer's own WhatsApp number is used for notifications and campaigns.",
        "Customer can capture leads from WhatsApp inbound messages, Instagram comments, Instagram DMs, and Meta sources.",
        "Customer can automate comment to DM, DM replies, lead follow-up, task creation, and campaign journeys.",
        "Partner can still control onboarding, customer mapping, API integration, and notification triggers from their own SaaS.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("Simple Pitch Line", level=1)
    doc.add_paragraph(
        "Marketing OS gives SaaS partners two things: API features for their own software, and a full marketing system "
        "their customers can use directly. Partners can onboard customers, connect each customer's WhatsApp number, "
        "send notifications from that number, and offer advanced marketing features like templates, broadcasts, flows, "
        "automation, WhatsApp leads, Instagram automation, Meta lead capture, analytics, campaigns, carousel, and catalog management."
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_docx()
